"""产出（v0.21）：把内容变成 .docx / .xlsx / .pptx / .pdf，以及纯文本类的
.md / .txt / .csv / .html（后者见 :func:`build_text`，不需要任何库）。

**方案 A：纯 Python 库，不装 LibreOffice / Pandoc。** 四条理由：

1. **读的那一半不用在这里。** 入库管线已经能解析这四种格式：docx / pptx 走
   `parsers/local_office.py`（纯本地，不需要 token），xlsx / csv 走
   `TabularParser`（带结构化副本），pdf 走 `parsers/local_pdf.py`（文字型）或云端
   MinerU（扫描件）。所以"能不能读懂这四种文件"这个问题早就答完了，
   缺的是另一半——**产出**。
2. **agent 的交付物得能发给别人。** 对话里整理出的一张表、一份报告，停在
   Markdown 上就只能在库里看；变成 .xlsx / .docx 才能被交给下游。
   而对方**点名要 .md / .txt 时，原样给出去就是交付**——换格式不是交付的
   必要条件（v0.41 补的 ``PLAIN_TEXT_KINDS`` 就是这一句的落点）。
3. **不装外部程序**：LibreOffice 是几百 MB 的安装包，Pandoc + Poppler 还要另外
   配路径；而它们能做到的事里，我们真正需要的（读写这四种格式）纯 Python 库都能做。
4. **代价写在明处**：这套东西只保证"结构正确、内容完整、文件能被
   Word / Excel / PowerPoint / 阅读器打开"，**不保证版式好看**。
   "渲染出来看一眼再改"需要一个真的排版引擎，方案 A 里没有它——
   所以本模块**不做任何"版式检查"**：没有渲染器时那种检查只能是假的
   （它要么恒真、要么在猜）。需要真排版时是方案 B（装 LibreOffice），
   那是一次显式的部署选择，不是这里偷偷假装的事。

**依赖是可选的**（与 parsers 那批同一个处置）：没装那个库时，
对应用户能读懂的一句话由 :func:`missing_requirement` 给出，
而不是抛一个 `ModuleNotFoundError` 到工具层变成"服务内部错误"。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

__all__ = [
    "MAX_CHARS",
    "MAX_ROWS",
    "MAX_SLIDES",
    "PLAIN_TEXT_KINDS",
    "build_docx",
    "build_pdf",
    "build_pptx",
    "build_text",
    "build_xlsx",
    "missing_requirement",
    "parse_blocks",
]

#: 正文上限。工具的输入来自模型，而"模型给了一段 50 万字的正文"这种事会发生。
MAX_CHARS = 200_000
#: 表格上限（行 × 列）。超了不是拒绝，是**说清楚哪一条超了**。
MAX_ROWS = 5_000
MAX_COLUMNS = 100
#: 幻灯片上限。一页一页写出来的是沟通材料，不是归档材料。
MAX_SLIDES = 60
MAX_BULLETS_PER_SLIDE = 20

#: **纯文本类产出**：正文原样落字节，不做任何转换，也不需要任何库。
#:
#: 为什么这四种也要走到这里来：交付口此前只认 .docx / .pdf（要过下面的转换器），
#: 于是"给我一份 .md / .txt"这种再普通不过的要求只能被回绝——模型实测的答复是
#: "导出文件只有那四种格式，没有 .md；沙箱也没开"，最后让用户自己复制。
#: **交付这件事不该挑格式**：.md 与 .docx 的差别只在对方拿它干什么，
#: 而不在"我们这边能不能造出来"。
PLAIN_TEXT_KINDS = ("md", "txt", "csv", "html")

#: 需要哪个库来完成哪种产出。键是"给用户看的名字"，值是 pip 包名。
_REQUIREMENTS = {
    "docx": ("python-docx", "docx"),
    "xlsx": ("openpyxl", "openpyxl"),
    "pptx": ("python-pptx", "pptx"),
    "pdf": ("reportlab", "reportlab"),
}


def missing_requirement(kind: str) -> str:
    """这种产出缺哪个依赖；齐了就回空串。

    **报"缺什么、怎么装"而不是抛 ImportError**：这个模块被工具层调用，
    而工具层的异常会被翻成给模型读的一句话——"服务内部错误"对它没有任何用，
    它需要的是"这件事现在做不到、原因是这个"。
    """
    spec = _REQUIREMENTS.get(kind)
    if spec is None:
        # 纯文本类不是"缺依赖"，是**根本没有依赖**——把它们报成"不认识的格式"，
        # 工具层就会在明明做得到的事情上回一句"做不到"
        return "" if kind in PLAIN_TEXT_KINDS else f"不认识的产出格式：{kind}"
    package, module = spec
    try:
        __import__(module)
    except ImportError:
        return (
            f"生成 {kind} 需要后端安装 {package}（`uv sync --extra office` 或 "
            f"`pip install {package}`），当前环境没有装"
        )
    return ""


# ------------------------------------------------------------------ 内容块


@dataclass(frozen=True, slots=True)
class Block:
    """正文里的一个块。**刻意只有四种**：标题、段落、列表项、表格。"""

    kind: str  # heading | paragraph | bullet | numbered | table
    text: str = ""
    level: int = 0
    rows: tuple[tuple[str, ...], ...] = ()


_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")


def parse_blocks(markdown: str) -> list[Block]:
    """把一段 Markdown 拆成块（**只认常见的那几种**）。

    为什么不引入一个 Markdown 库：我们要的是"能放进 docx/pdf 的结构"，
    而这个子集（标题、段落、列表、表格）覆盖了 95% 的实际输入；
    引一个完整实现反而会带进一堆我们表达不了的东西（脚注、内嵌 HTML、
    任务列表），那时只能丢掉——**丢得静悄悄还不如一开始就不认**。
    """
    blocks: list[Block] = []
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        heading = _HEADING.match(stripped)
        if heading:
            blocks.append(Block("heading", heading.group(2).strip(), level=len(heading.group(1))))
            index += 1
            continue

        if _TABLE_ROW.match(stripped):
            rows: list[tuple[str, ...]] = []
            while index < len(lines) and _TABLE_ROW.match(lines[index].strip()):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                # 分隔行（`| --- | --- |`）是 Markdown 表格的语法，不是数据
                if not all(re.fullmatch(r":?-{2,}:?", cell or "-") for cell in cells):
                    rows.append(tuple(cells))
                index += 1
            if rows:
                blocks.append(Block("table", rows=tuple(rows)))
            continue

        bullet = _BULLET.match(line)
        if bullet:
            blocks.append(Block("bullet", bullet.group(1).strip()))
            index += 1
            continue

        numbered = _NUMBERED.match(line)
        if numbered:
            blocks.append(Block("numbered", numbered.group(1).strip()))
            index += 1
            continue

        # 剩下的按段落接：连续的非空行合成一段（Markdown 里那本来就是一个段落）
        paragraph: list[str] = []
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or _HEADING.match(candidate) or _BULLET.match(lines[index]):
                break
            if _TABLE_ROW.match(candidate) or _NUMBERED.match(lines[index]):
                break
            paragraph.append(candidate)
            index += 1
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
    return blocks


# ------------------------------------------------------------------ 纯文本


def build_text(text: str, *, kind: str) -> bytes:
    """正文 → .md / .txt / .csv / .html（**原样落字节**，见 :data:`PLAIN_TEXT_KINDS`）。

    **不做 Markdown → 纯文本的改写**：模型给的正文就是对方要的那份东西，
    我们再去一遍标记（``# 标题`` → ``标题``、``| a | b |`` → 一行竖线）
    只会让他下载到的东西与屏幕上看到的不是同一份——而"我拿到的和你给我看的不一样"
    是最没法解释的一类问题。

    ``title`` 不参与：纯文本产出的正文就是全文，而文件名已经带着标题——
    再往正文里插一个标题，等于往对方要的那份内容里加东西（工具层因此只为
    .docx / .pdf 那一支传它）。

    ``.csv`` 是唯一一处加工：**加一个 UTF-8 BOM**。Excel / WPS 双击打开无 BOM 的
    UTF-8 CSV 会按本地编码（中文环境是 GBK）解，中文全是乱码——而一份 .csv 的去向
    几乎总是"被 Excel 打开"。BOM 对别的读法无害：我们自己的解析器就是按
    ``utf-8-sig`` 起头解的（见 `parsers/text_decode.py` 的解码阶梯）。
    """
    if kind not in PLAIN_TEXT_KINDS:
        raise RuntimeError(
            f"{kind} 不是纯文本产出（只做 "
            f"{'、'.join('.' + item for item in PLAIN_TEXT_KINDS)}）"
        )
    body = (text or "").encode("utf-8")
    if kind == "csv":
        return b"\xef\xbb\xbf" + body
    return body


# ------------------------------------------------------------------ docx


def build_docx(markdown: str, *, title: str = "") -> bytes:
    """Markdown → .docx。表格用真表格，标题用真标题样式。"""
    problem = missing_requirement("docx")
    if problem:
        raise RuntimeError(problem)

    from docx import Document
    from docx.shared import Pt

    document = Document()
    # 中文字体要显式指到 eastAsia 上：只设 `font.name` 管的是西文，
    # 中文会回落到 Word 的默认值——而不同机器上的默认值不一样，
    # 同一份文件在两台机器上排出来的行数就不同。
    style = document.styles["Normal"]
    style.font.size = Pt(11)
    _set_east_asia(style, "宋体")

    if title.strip():
        document.add_heading(title.strip(), level=0)

    for block in parse_blocks(markdown):
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level, 4))
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "numbered":
            document.add_paragraph(block.text, style="List Number")
        elif block.kind == "table":
            _docx_table(document, block)
        else:
            document.add_paragraph(block.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _set_east_asia(style: Any, font_name: str) -> None:
    """把样式的中文字体指到 `w:eastAsia`（python-docx 没有直接的 API）。"""
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def _docx_table(document: Any, block: Block) -> None:
    rows = block.rows
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            table.cell(r, c).text = row[c] if c < len(row) else ""


# ------------------------------------------------------------------ xlsx


def build_xlsx(rows: list[list[Any]], *, sheet_name: str = "Sheet1") -> bytes:
    """二维数据 → .xlsx。**只写值**：数值写数值、其余写文本。

    数值要按数值写（而不是一律转成字符串）：Excel 里的求和、排序、图表都依赖
    单元格类型，全写成文本的话它们全部失效——而用户会以为是我们算错了。
    """
    problem = missing_requirement("xlsx")
    if problem:
        raise RuntimeError(problem)

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name or "Sheet1")[:31]  # Excel 的表名上限就是 31 字

    widths: dict[int, int] = {}
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            sheet.cell(row=r, column=c, value=_cell_value(value))
            widths[c] = max(widths.get(c, 6), min(40, len(str(value)) + 2))

    # 列宽跟着内容走：不设的话打开是"一列挤成一堆 ####"，每次都要手动拉
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    # 纯数字的字符串按数字写：模型从正文里抠出来的数字常常是字符串，
    # 而"一列数字全是文本"是 Excel 里最常见的坑
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)
    return text


# ------------------------------------------------------------------ pptx


def build_pptx(slides: list[dict[str, Any]], *, title: str = "") -> bytes:
    """幻灯片 → .pptx。``slides`` 是 ``[{"title": …, "bullets": […]}]``。

    **一页一件事**：正文只收要点（bullets），不收整段——把一整段塞进一页
    是"用 PPT 当 Word"最常见的做法，而它的结果是既不好讲也不好读。
    需要成段文字时那份材料适合 docx（见 :func:`build_docx`）。
    """
    problem = missing_requirement("pptx")
    if problem:
        raise RuntimeError(problem)

    from pptx import Presentation
    from pptx.util import Pt

    deck = Presentation()
    if title.strip():
        cover = deck.slides.add_slide(deck.slide_layouts[0])
        cover.shapes.title.text = title.strip()

    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title") or "")
        body = slide.placeholders[1].text_frame
        body.clear()
        for position, bullet in enumerate(item.get("bullets") or []):
            paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(18)

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ pdf


def build_pdf(markdown: str, *, title: str = "") -> bytes:
    """Markdown → .pdf（reportlab 的 Platypus，自动分页）。

    **中文必须显式注册一个 CJK 字体**：reportlab 内置的 Helvetica 不含汉字，
    直接用会得到一页黑方块（而且不报错——它认为那是正常输出）。
    这里用 reportlab 自带的 CID 字体 ``STSong-Light``：不需要额外的字体文件，
    阅读器那边按标准 CJK 字体名渲染。
    """
    problem = missing_requirement("pdf")
    if problem:
        raise RuntimeError(problem)

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "KylabBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10.5, leading=16
    )
    heading_styles = {
        level: ParagraphStyle(
            f"KylabH{level}",
            parent=styles["Heading1"],
            fontName="STSong-Light",
            fontSize=max(12, 18 - level * 2),
            leading=max(18, 24 - level * 2),
        )
        for level in range(1, 5)
    }

    story: list[Any] = []
    if title.strip():
        story.append(Paragraph(_escape(title.strip()), heading_styles[1]))
        story.append(Spacer(1, 4 * mm))
    for block in parse_blocks(markdown):
        if block.kind == "heading":
            story.append(Paragraph(_escape(block.text), heading_styles[min(block.level, 4)]))
        elif block.kind in ("bullet", "numbered"):
            prefix = "• " if block.kind == "bullet" else "1. "
            story.append(Paragraph(prefix + _escape(block.text), body))
        elif block.kind == "table":
            story.append(_pdf_table(block, Table, TableStyle))
        else:
            story.append(Paragraph(_escape(block.text), body))
        story.append(Spacer(1, 2 * mm))

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, title=title or "KYLAB", leftMargin=20 * mm, rightMargin=20 * mm
    ).build(story)
    return buffer.getvalue()


def _escape(text: str) -> str:
    """reportlab 的 Paragraph 认一小套标记语言，正文里的 ``&`` ``<`` 必须先转义。"""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_table(block: Block, table_cls: Any, style_cls: Any) -> Any:
    rows = [[_escape(cell) for cell in row] for row in block.rows]
    shape = table_cls(rows)
    shape.setStyle(
        style_cls(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, (0.6, 0.6, 0.6)),
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return shape
